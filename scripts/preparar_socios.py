#!/usr/bin/env python3
"""
Convierte la lista de asociados (Excel o CSV) en un archivo SQL listo para cargar a la base de datos.

USO:
    python scripts/preparar_socios.py  ruta/a/asociados.xlsx
    python scripts/preparar_socios.py  ruta/a/asociados.csv

Genera, en la misma carpeta del archivo original:
    socios-importar.sql     -> lo que se carga a la base de datos
                               (si hay más de 100 asociados se divide en partes: socios-importar-1.sql,
                                socios-importar-2.sql… y se cargan una por una, en orden)
    socios-revision.txt     -> filas con problemas (solo número de fila y motivo, sin datos personales)

Columnas que entiende (no importa el orden, ni las mayúsculas, ni las tildes):
    número de socio · nombre · email/correo · teléfono · fecha de nacimiento · sector
La columna de IDENTIFICACIÓN / CÉDULA se IGNORA a propósito: el acceso no la necesita y
mientras menos datos sensibles guardemos, menos riesgo hay.

Es seguro correrlo varias veces: si un asociado ya existe, se actualizan sus datos
(nombre, correo, teléfono, nacimiento, sector) sin tocar su estado ni su cuenta.

IMPORTANTE: estos archivos tienen datos personales. NO los subas a GitHub
(el .gitignore ya los excluye si los dejas con nombres que empiecen con "socios").
"""
import csv
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

COLUMNAS = {
    "numero_socio": {"numerodesocio", "numerosocio", "nsocio", "nrosocio", "nosocio", "socio", "numero", "codigosocio"},
    "nombre": {"nombre", "nombrecompleto", "nombreyapellidos", "asociado", "nombredelcliente", "cliente", "nombredelsocio", "nombredelasociado"},
    "correo": {"email", "correo", "correoelectronico", "mail"},
    "telefono": {"telefono", "celular", "tel", "telefonocelular"},
    "fecha_nacimiento": {"fechadenacimiento", "fechanacimiento", "nacimiento", "fechanac", "cumpleanos"},
    "sector": {"sector", "zona"},
}
IGNORADAS = {"identificacion", "cedula", "id", "dni", "numerodeidentificacion", "cedulaidentidad"}
FORMATOS_FECHA = ["%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y"]


def clave(texto):
    t = unicodedata.normalize("NFD", str(texto))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", t.lower())


def leer_filas(ruta):
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Para leer .xlsx instala openpyxl (pip install openpyxl) o guarda el Excel como CSV UTF-8.")
        hoja = load_workbook(ruta, read_only=True, data_only=True).worksheets[0]
        return [list(f) for f in hoja.iter_rows(values_only=True)]
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = ruta.read_text(encoding=codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:4000]
    delimitador = max([",", ";", "\t"], key=muestra.count)
    return list(csv.reader(texto.splitlines(), delimiter=delimitador))


def parse_fecha(valor):
    if valor is None or str(valor).strip() == "":
        return None, None
    if isinstance(valor, (datetime, date)):
        d = valor.date() if isinstance(valor, datetime) else valor
    else:
        d = None
        for f in FORMATOS_FECHA:
            try:
                d = datetime.strptime(str(valor).strip().split(" ")[0], f).date()
                break
            except ValueError:
                continue
        if d is None:
            return None, f"fecha de nacimiento no entendida"
        if d.year > date.today().year:  # años de 2 dígitos: 85 -> 1985, no 2085
            d = d.replace(year=d.year - 100)
    if not (1900 <= d.year <= date.today().year):
        return None, "fecha de nacimiento fuera de rango"
    return d.isoformat(), None


def limpiar_numero(valor):
    t = str(valor).strip()
    if re.fullmatch(r"\d+\.0", t):  # Excel guarda 1001 como 1001.0
        t = t[:-2]
    return re.sub(r"\s+", "", t)[:30]


def sql(valor):
    if valor is None or valor == "":
        return "NULL"
    return "'" + str(valor).replace("'", "''") + "'"


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    ruta = Path(sys.argv[1])
    if not ruta.exists():
        sys.exit(f"No encuentro el archivo: {ruta}")

    filas = leer_filas(ruta)
    filas = [f for f in filas if any(str(c or "").strip() for c in f)]
    if not filas:
        sys.exit("El archivo está vacío.")

    encabezado = [clave(c) for c in filas[0]]
    indice, ignoradas = {}, []
    for i, nombre in enumerate(encabezado):
        for campo, alias in COLUMNAS.items():
            if nombre in alias and campo not in indice:
                indice[campo] = i
        if nombre in IGNORADAS:
            ignoradas.append(str(filas[0][i]))
    # Segunda pasada: palabras clave dentro del encabezado (para títulos como "Nombre del Cliente")
    PALABRAS = [
        ("nombre", ["nombre", "cliente"]),
        ("correo", ["mail", "correo"]),
        ("telefono", ["telefono", "celular"]),
        ("fecha_nacimiento", ["nacimiento", "cumple"]),
        ("sector", ["sector"]),
        ("numero_socio", ["socio"]),
    ]
    usadas = set(indice.values())
    for campo, claves in PALABRAS:
        if campo in indice:
            continue
        for i, nombre in enumerate(encabezado):
            if i in usadas or nombre in IGNORADAS:
                continue
            if any(k in nombre for k in claves):
                indice[campo] = i
                usadas.add(i)
                break

    faltan = [c for c in ("numero_socio", "nombre", "correo") if c not in indice]
    if faltan:
        sys.exit(f"No encontré estas columnas obligatorias: {', '.join(faltan)}.\nEncabezados leídos: {[str(c) for c in filas[0]]}")

    print("Columnas usadas:")
    for campo, i in indice.items():
        print(f"   {campo:17} <- \"{filas[0][i]}\"")

    def celda(fila, campo):
        i = indice.get(campo)
        return fila[i] if i is not None and i < len(fila) else None

    ok, problemas, vistos, correos = [], [], set(), {}
    for n, fila in enumerate(filas[1:], start=2):  # n = número de fila en Excel
        numero = limpiar_numero(celda(fila, "numero_socio") or "")
        nombre = " ".join(str(celda(fila, "nombre") or "").split())
        correo = str(celda(fila, "correo") or "").strip().lower()
        if not numero or not nombre:
            problemas.append((n, "falta número de socio o nombre: no se importó"))
            continue
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", correo):
            problemas.append((n, "correo vacío o inválido: no se importó (esa persona no podría entrar)"))
            continue
        if numero in vistos:
            problemas.append((n, "número de socio repetido en el archivo: se conservó la primera fila"))
            continue
        vistos.add(numero)
        nacimiento, aviso = parse_fecha(celda(fila, "fecha_nacimiento"))
        if aviso:
            problemas.append((n, f"{aviso}: se importó sin fecha"))
        telefono = re.sub(r"[^\d+]", "", str(celda(fila, "telefono") or "")) or None
        sector = " ".join(str(celda(fila, "sector") or "").split()) or None
        ok.append((numero, nombre, correo, telefono, nacimiento, sector))
        correos[correo] = correos.get(correo, 0) + 1

    encabezado_sql = "-- Generado por scripts/preparar_socios.py. Contiene datos personales: NO subir a GitHub."
    inserts = []
    for numero, nombre, correo, telefono, nacimiento, sector in ok:
        inserts.append(
            "INSERT INTO socios (numero_socio, nombre, correo, telefono, fecha_nacimiento, sector) VALUES "
            f"({sql(numero)}, {sql(nombre)}, {sql(correo)}, {sql(telefono)}, {sql(nacimiento)}, {sql(sector)}) "
            "ON CONFLICT(numero_socio) DO UPDATE SET nombre = excluded.nombre, correo = excluded.correo, "
            "telefono = excluded.telefono, fecha_nacimiento = excluded.fecha_nacimiento, sector = excluded.sector;"
        )

    # Se borran los archivos de corridas anteriores para no cargar uno viejo por error
    for viejo in ruta.parent.glob("socios-importar*.sql"):
        viejo.unlink()
    TAMANO_PARTE = 100  # asociados por archivo: la consola de D1 se queja si el texto es muy largo
    partes = [inserts[i:i + TAMANO_PARTE] for i in range(0, len(inserts), TAMANO_PARTE)] or [[]]
    archivos_sql = []
    for n, parte in enumerate(partes, start=1):
        nombre_archivo = "socios-importar.sql" if len(partes) == 1 else f"socios-importar-{n}.sql"
        destino = ruta.with_name(nombre_archivo)
        destino.write_text(encabezado_sql + "\n" + "\n".join(parte) + "\n", encoding="utf-8")
        archivos_sql.append(destino)

    compartidos = sum(1 for c in correos.values() if c > 1)
    reporte = ruta.with_name("socios-revision.txt")
    texto = [f"Filas leídas: {len(filas) - 1} | listas para importar: {len(ok)} | con problemas: {len(problemas)}", ""]
    texto += [f"Fila {n}: {motivo}" for n, motivo in problemas]
    if compartidos:
        texto += ["", f"Aviso: {compartidos} correo(s) están en más de un asociado (p. ej. familias). Pueden entrar, porque el acceso pide número de socio + correo."]
    reporte.write_text("\n".join(texto) + "\n", encoding="utf-8")

    print(f"Listos para importar: {len(ok)} de {len(filas) - 1}")
    if ignoradas:
        print(f"Columna(s) ignorada(s) a propósito (no se guardan): {', '.join(ignoradas)}")
    if problemas:
        print(f"Con problemas: {len(problemas)}  -> revisa {reporte.name}")
    print("Archivo(s) para cargar a la base de datos, en este orden:")
    for a in archivos_sql:
        print(f"   {a}")


if __name__ == "__main__":
    main()
